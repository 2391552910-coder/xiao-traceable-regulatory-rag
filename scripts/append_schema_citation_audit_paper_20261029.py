"""Add mechanical citation-ID audit results to existing working drafts."""
from pathlib import Path
from docx import Document
from docx.shared import Inches

root = Path(__file__).resolve().parents[1]
items = [
    ('results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_格式提示补充候选稿_待终审.docx',
     'results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_引用编号审计候选稿_待终审.docx',
     '事后格式提示开发试验。',
     ' 对引用的机械审计进一步发现：示例提示下仅 2/8 条回答的三个字段均有非空、可在输入材料中找到的有效编号；另有一条虽满足字段类型，却把日期写入引用列表。此项检查不判断所引材料是否支持具体内容。'),
    ('results/paper/english_translation_current_20261028_01/Controlled_Evaluation_Dual_Source_RAG_schema_prompt_development_20261029.docx',
     'results/paper/english_translation_current_20261028_01/Controlled_Evaluation_Dual_Source_RAG_citation_id_audit_20261029.docx',
     'Post-hoc format-prompt development.',
     ' A further mechanical citation audit found that only 2/8 example-prompt answers provided nonempty, resolvable source IDs for all three fields; another answer met the type checks but put date strings in its citation lists. This audit does not determine whether the cited material supports the specific answer.')]
for source, output, start, addition in items:
    document = Document(root/source)
    matches = [p for p in document.paragraphs if p.text.startswith(start)]
    if len(matches) != 1:
        raise ValueError('Expected one prompt-development paragraph')
    matches[0].add_run(addition)
    if start.startswith('Post-hoc'):
        image = document.inline_shapes[-1]
        image.height = int(image.height * Inches(4.3) / image.width)
        image.width = Inches(4.3)
    document.save(root/output)
    print(root/output)
