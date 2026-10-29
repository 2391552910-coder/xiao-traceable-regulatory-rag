"""Add a transparent post-hoc prompt development note to existing drafts."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root = Path(__file__).resolve().parents[1]
image = root / 'results/runs/qwen_schema_prompt_figures_20261029_01/schema_prompt_development.png'
items = [
    (next((root/'results/paper/condensed_candidate_20261028_01').glob('*跨模型与Agent*.docx')),
     root/'results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_格式提示补充候选稿_待终审.docx',
     '事后格式提示开发试验。针对观察到的 Qwen 引用字段格式错误，在原提示末尾增加空结构示例，并以相同模型、八道题、证据、采样设置和种子重新生成。按事后勘误的既有助手提取参考，结构合规由 0/8 增至 3/8，成员宏平均 F1 由 0.651 至 0.678，日期宏平均 F1 由 0.600 降至 0.558，会议次数完全匹配由 3/8 至 6/8，三项同时完全匹配均为 1/8。原提示与示例提示分别使用输入 38998 和 39710 tokens、输出 2695 和 1114 tokens。该提示根据同一批已观察错误设计，属于事后开发而非独立测试；一次采样和小样本不能建立稳健改进，结构合规也不能证明引用支持答案。既有参考未经独立验证，本试验未新增人工标注。',
     '图 11  同一模型的事后格式提示开发试验。使用既有助手提取并事后勘误的参考；结构合规不等于引用语义正确。'),
    (next((root/'results/paper/english_translation_current_20261028_01').glob('*cross_model_agent*.docx')),
     root/'results/paper/english_translation_current_20261028_01/Controlled_Evaluation_Dual_Source_RAG_schema_prompt_development_20261029.docx',
     'Post-hoc format-prompt development. After inspecting Qwen citation-field errors, we appended an empty schema example and regenerated answers for the same eight questions and evidence with the same model, sampling settings and seeds. Against the existing assistant-extracted references with a post-output erratum, schema compliance rose from 0/8 to 3/8, member macro F1 from 0.651 to 0.678, date macro F1 fell from 0.600 to 0.558, and exact meeting-count matches rose from 3/8 to 6/8; all three fields matched exactly on 1/8 questions in both conditions. The original and example prompts used 38,998 and 39,710 input tokens and produced 2,695 and 1,114 output tokens, respectively. This prompt was designed after observing errors on these same cases and is a development experiment, not independent validation. A single sample and small dataset cannot establish a robust improvement, and schema validity does not establish semantic citation support. The existing references have not been independently validated; no new human labels were collected.',
     'Figure 11. Post-hoc within-model format-prompt development using existing assistant-extracted references with a post-output erratum. Schema validity does not establish citation support.')]

for source, target, note, caption in items:
    d = Document(source)
    anchor = next(p for p in d.paragraphs if p.text.startswith('7.7 '))
    def insert(style=None):
        element = OxmlElement('w:p'); anchor._p.addprevious(element)
        paragraph = Paragraph(element, anchor._parent)
        if style: paragraph.style = style
        return paragraph
    text = insert(); text.add_run(note)
    fig = insert(); fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.add_run().add_picture(str(image), width=Inches(5.0))
    cap = insert(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.add_run(caption)
    d.save(target)
    print(target)
