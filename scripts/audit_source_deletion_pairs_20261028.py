from __future__ import annotations
import hashlib
import json
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCORES = ROOT / "deliverables/20261028_provider_corrected_01/analysis/per_output_scores_primary_model_only.json"
MANUSCRIPT = ROOT / "results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_统一主文_来源缺失配对口径修订候选稿_待终审.docx"
OUTPUT = ROOT / "results/paper/condensed_candidate_20261028_01/source_deletion_pair_audit_20261028.json"
G3 = "dual_source_role_grounded_rag"
CONDITIONS = {"remove_annual_report": "minus_annual", "remove_regulation": "minus_regulation"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def normalized_paragraph_hash(docx_path: Path, index: int) -> str:
    from docx import Document
    text = " ".join(Document(docx_path).paragraphs[index].text.split())
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

rows = json.loads(SCORES.read_text(encoding="utf-8"))
by_question: dict[str, dict[str, dict]] = defaultdict(dict)
for row in rows:
    by_question[row["question_id"]][row["mode"]] = row

results = {}
for label, condition in CONDITIONS.items():
    condition_rows = [r for r in rows if r["mode"] == condition]
    paired = []
    for qid, conditions in by_question.items():
        if G3 in conditions and condition in conditions:
            paired.append((qid, conditions[G3], conditions[condition]))
    diffs = [g3["mean"]["total"] - removed["mean"]["total"] for _, g3, removed in paired]
    qids_with_removed = [qid for qid, conditions in by_question.items() if condition in conditions]
    missing_g3 = [qid for qid in qids_with_removed if G3 not in by_question[qid]]
    known_diff_sum = sum(diff for diff in diffs)
    removed_missing_sum = sum(by_question[qid][condition]["mean"]["total"] for qid in missing_g3)
    n_planned = len(qids_with_removed)
    bounds = [
        (known_diff_sum - removed_missing_sum) / n_planned,
        (known_diff_sum + 8 * len(missing_g3) - removed_missing_sum) / n_planned,
    ]
    results[label] = {
        "condition": condition,
        "valid_outputs": len(condition_rows),
        "paired_questions": len(paired),
        "paired_companies": len({removed["company"] for _, _, removed in paired}),
        "paired_mean_difference_g3_minus_removed": sum(diffs) / len(diffs),
        "paired_difference_min": min(diffs),
        "paired_difference_max": max(diffs),
        "planned_removal_questions": n_planned,
        "missing_g3_outputs_in_removal_questions": len(missing_g3),
        "sharp_mean_difference_bounds_over_removal_question_set": bounds,
        "bounds_assumption": "G3 missing outcomes are bounded only by the observed 0-8 rubric; all source-removal outcomes are observed.",
    }

from docx import Document
paragraph = " ".join(Document(MANUSCRIPT).paragraphs[139].text.split())
report = {
    "audit_date": "2026-10-28",
    "scope": "Paired reanalysis of the existing source-removal outputs; no new model calls or labels.",
    "manuscript": {
        "path": str(MANUSCRIPT.relative_to(ROOT)).replace("\\", "/"),
        "sha256": sha256(MANUSCRIPT),
        "paragraph_index_zero_based": 139,
        "paragraph_sha256_utf8": hashlib.sha256(paragraph.encode("utf-8")).hexdigest(),
    },
    "scores": {
        "path": str(SCORES.relative_to(ROOT)).replace("\\", "/"),
        "sha256": sha256(SCORES),
        "row_count": len(rows),
    },
    "comparison": "G3 minus the same-question source-removal condition; total score is the two-reviewer mean on the 0-8 rubric. Each paired question is from a different company, so the paired mean also equals the company-macro mean.",
    "results": results,
    "interpretation": "Observed paired differences are conditional on 15 available G3 pairs. The sharp bounds over all 20 planned source-removal questions include zero. Removing a source also changes evidence content and context length, so these results do not isolate a pure source-role effect.",
}
OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(OUTPUT)
print(json.dumps(report["results"], ensure_ascii=False, indent=2))

