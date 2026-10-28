"""Create a corrected manuscript candidate using provider outputs only."""
from __future__ import annotations
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json, os, tempfile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_精简主文候选稿_待终审.docx"
OUTDIR = ROOT / "results/paper/primary_provider_candidate_20261028_01"
OUT = OUTDIR / "年报与法规双源RAG_供应商输出重分析候选稿_待渲染验收.docx"
PLOTS = ROOT / "results/runs/primary_provider_only_robustness_20261028_01"


def set_text(paragraph, value):
    paragraph.text = value


def set_cell(cell, value):
    cell.text = value


def replace_paragraph(doc, marker, value, *, starts=False):
    matches = [p for p in doc.paragraphs if (p.text.startswith(marker) if starts else marker in p.text)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {marker!r}, got {len(matches)}")
    set_text(matches[0], value)


def update_tables(doc):
    # Corpus and scored-review scope.
    set_cell(doc.tables[3].rows[-1].cells[0], "全条件供应商输出评分")
    set_cell(doc.tables[3].rows[-1].cells[1], "355回答 × 2人 × 4维")
    set_cell(doc.tables[3].rows[-1].cells[2], "2840个0/1/2评分值；33条会话补充输出不纳入，其中31条曾进入旧版评分汇总")

    # Planned conditions and the corrected all-condition denominator.
    set_cell(doc.tables[4].rows[1].cells[2], "G1/G2/G3；仅计冻结供应商调用输出")
    set_cell(doc.tables[4].rows[6].cells[0], "供应商输出评分与稳健性诊断")
    set_cell(doc.tables[4].rows[6].cells[1], "355输出 / 400计划单元")
    set_cell(doc.tables[4].rows[6].cells[2], "排除33条会话补充回答；355条供应商输出由两人评分")

    # All condition summaries, using the two-reviewer mean of eligible provider outputs.
    table = doc.tables[5]
    replacements = {
        "G1": ("52/60", "7.327", "7.577", "7.077"),
        "G2": ("47/60", "7.543", "7.809", "7.277"),
        "G3": ("48/60", "7.594", "7.833", "7.354"),
    }
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        if label in replacements:
            for cell, value in zip(row.cells[1:], replacements[label]):
                set_cell(cell, value)

    # Paired effects on available provider-generated outputs.
    table = doc.tables[6]
    effect_rows = [
        ("G2-G1", "41/60", "0.228", "[0.110, 0.355]"),
        ("G3-G2", "38/60", "0.013", "[0.000, 0.039]"),
        ("G3-G1", "42/60", "0.289", "[0.158, 0.434]"),
        ("G3-NoRerank", "45/60", "-0.222", "[-0.315, -0.130]"),
        ("G3-BM25", "45/60", "-0.245", "[-0.343, -0.157]"),
    ]
    for row, values in zip(table.rows[1:], effect_rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)

    # Keep the main table focused on the three primary conditions.
    table = doc.tables[7]
    dimension_rows = [
        ("G2-G1", "41", "-0.009", "0.237", "0.000", "0.000", "0.228"),
        ("G3-G2", "38", "0.000", "0.000", "0.013", "0.000", "0.013"),
        ("G3-G1", "42", "-0.009", "0.285", "0.013", "0.000", "0.289"),
    ]
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for values in dimension_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell(cell, value)

    # Agreement and ceiling diagnostics; constant uncertainty marginals make kappa undefined.
    table = doc.tables[8]
    reliability_rows = [
        ("公司事实", "86.2%", "0.400", "0.116", "81.4%", "93.0%"),
        ("法规依据", "57.7%", "0.215", "-0.417", "86.5%", "44.8%"),
        ("引用支持", "97.5%", "0.000", "0.025", "97.5%", "100.0%"),
        ("不确定性", "100.0%", "不可估计†", "0.000", "100.0%", "100.0%"),
    ]
    for row, values in zip(table.rows[1:], reliability_rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)

    # Missingness and leave-one-company-out table.
    table = doc.tables[9]
    sensitivity_rows = [
        ("G2-G1", "[0.199, 0.241]", "[-1.508, 1.292]", "13/28/0"),
        ("G3-G2", "[0.000, 0.014]", "[-1.567, 1.767]", "1/37/0"),
        ("G3-NoRerank", "[-0.250, -0.196]", "[-1.750, 0.517]", "2/23/20"),
        ("G3-BM25", "[-0.260, -0.219]", "[-1.792, 0.608]", "10/6/29"),
    ]
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for values in sensitivity_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell(cell, value)

    # Replace contaminated post-hoc examples with a fully provider-generated case.
    table = doc.tables[11]
    first = table.rows[1]
    for row in list(table.rows[2:]):
        table._tbl.remove(row._tr)
    values = [
        "002475SZ-2024-AUDIT_COMMITTEE",
        "7.0/8.0/8.0",
        "一、年报披露（有原文支持）：1. 人员构成：立讯精密2024年报披露，审计委员会成员为刘中华、宋宇红、侯玲玲、张英，共4人；报告期内召开会议4次 [A1]。年报另披露公司董事会共7名董事，其中独立董事3名（法律、会计等领域专家），董事会下设审计委员会、战略委员会、提名委员会、薪酬与考核委员会4个专门委员会，称‘成员组成结构合理’ [A4]，但未逐一说明审计委员会各成员的董事类别及召集人身份……",
        "A1,A4,R1,R2,R3,R4,R5",
    ]
    for cell, value in zip(first.cells, values):
        set_cell(cell, value)


def replace_embedded_media(docx_path):
    replacements = {
        "word/media/image4.png": PLOTS / "primary_scores_provider_only.png",
        "word/media/image5.png": PLOTS / "dimension_effects_provider_only.png",
        "word/media/image7.png": PLOTS / "reviewer_ceiling_provider_only.png",
        "word/media/image8.png": PLOTS / "missingness_loco_provider_only.png",
    }
    for p in replacements.values():
        if not p.exists():
            raise FileNotFoundError(p)
    fd, tmp_name = tempfile.mkstemp(suffix=".docx", dir=OUTDIR)
    os.close(fd)
    try:
        with ZipFile(docx_path, "r") as src, ZipFile(tmp_name, "w", ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = replacements[item.filename].read_bytes() if item.filename in replacements else src.read(item.filename)
                dst.writestr(item, data)
        os.replace(tmp_name, docx_path)
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    if OUT.exists():
        raise ValueError(f"Refusing to overwrite an existing manuscript candidate: {OUT}")
    if any(OUTDIR.iterdir()):
        raise ValueError(f"Output directory already contains files: {OUTDIR}")
    doc = Document(SOURCE)
    replace_paragraph(doc, "上市公司评价既需要", "上市公司评价既需要年报披露，也需要相应法规依据。本文考察在年报与法规问答中，按来源分路检索相较合并检索能否改善回答质量及证据供给。我们构建保留文档身份和原文坐标的双源检索增强生成流程，在20家公司2024年年报的60个问题上比较合并检索、双路检索及相同证据下的来源角色提示。主对照原计划180个Kimi供应商请求，首轮147个结构有效；33条会话补充回答未执行冻结提示及API协议，其中31条曾进入旧评分汇总，现全部排除。过滤后保留355/400个供应商生成输出的双人评分。G2相较G1在41个可配对问题上的0至8分公司宏均差为0.228（20,000次公司聚类bootstrap点对点95%区间[0.110, 0.355]），但全60题的缺失数据界跨0。G3相较G2在38个配对问题上的差为0.013（区间[0.000, 0.039]）。双路检索使两类来源同现率由68.3%升至100%，但平均参考位置覆盖未提高，法规参考位置覆盖仍有限。观察到的总分差异主要集中在法规评分维度；引用及不确定性维度存在天花板。结果仅适用于本研究样本、模型和评分协议。")
    replace_paragraph(doc, "Listed-company evaluation requires", "Listed-company evaluation requires both annual-report disclosures and relevant regulatory evidence. This study asks whether source-separated retrieval improves answer ratings and evidence supply over pooled retrieval. We build a traceable dual-source RAG pipeline and compare merged retrieval, dual retrieval, and source-role prompting with identical evidence on 60 questions from the 2024 annual reports of 20 companies. The primary comparison planned 180 Kimi-provider requests; 147 were structurally valid in the first run. Thirty-three session-authored supplements did not execute the frozen prompt or API protocol; 31 had entered the earlier score summary and are now excluded. The filtered set contains 355 two-reviewer-scored provider outputs from 400 planned units. On 41 paired questions, dual retrieval exceeds pooled retrieval by 0.228 points on a company-macro 0–8 scale (20,000-draw company-clustered pointwise 95% interval: [0.110, 0.355]); however, the missing-outcome bounds for all 60 questions include zero. Source-role prompting adds 0.013 points on 38 observed pairs (interval: [0.000, 0.039]). Dual retrieval increases two-source presence from 68.3% to 100%, while mean reference-location coverage does not improve and regulatory reference coverage remains limited. The observed score difference is concentrated in the regulation dimension; citation and uncertainty ratings show ceiling effects. These findings are limited to the sampled companies, model configuration, observed pairs, and rating protocol.")

    replace_paragraph(doc, "完整原始响应、结构失败与一次恢复视图", "完整原始响应、结构失败及会话补充记录均保留；会话补充没有新增供应商API调用，未纳入受控评分分析。新增分析使用新目录，不覆盖原评分、运行或历史统计版本。")
    replace_paragraph(doc, "主生成temperature为1", "主生成temperature为1，最大输出16384 tokens，请求超时600秒，SDK自动重试为0，每个题目—条件执行一次主请求。调度随机种子42没有发送为生成随机种子，不能据此宣称可精确复现采样。准备阶段因供应商参数限制或截断发生的请求设置修订及错误回执均保存。首轮180个主对照请求中147个结构有效、33个失败。随后保存的33条补充回答由当前助手会话撰写，未执行冻结提示，新增供应商API调用为0；31条曾被双人评分，2条未评分。因生成来源和协议不同，33条均从受控比较及全条件统计中排除。")
    replace_paragraph(doc, "两位人员分别对实际回答核对原文", "两位人员分别对实际回答核对原文。主对照仅纳入147条首轮供应商API结构有效输出；所有实验条件过滤后共保留355条供应商输出的双人评分。评分包含公司事实、法规依据、引用支持和不确定性处理，每维0、1、2或U。只有两人四维均为数值时计算均分总量Q_i=(1/2)∑_r∑_j s_irj，范围0至8。本轮纳入评分的输出没有U；失败与未运行不填质量0分。均分是分析指标，不是分歧裁定或一致金标。")

    replace_paragraph(doc, "400个计划回答中，主对照恢复视图", "400个计划供应商请求中，主对照首轮147个、检索消融109个、来源缺失40个、同系列第二版本59个，共355个结构有效输出且取得双人评分。另有33条会话补充回答，其中31条被旧评分清单收录；它们未执行冻结的Kimi提示和API协议，现已排除。排除补充回答后，45个计划供应商单元没有合格输出。有效表示供应商返回且通过结构检查，不表示事实、引用或法规适用性正确。")
    replace_paragraph(doc, "主对照G2−G1的公司宏差值为", "在纳入的供应商输出中，G2−G1的41个已知配对问题公司宏差值为0.228，20,000次公司聚类bootstrap点对点95%区间为[0.110, 0.355]；G3−G1为0.289（42对，[0.158, 0.434]）。G3−G2仅有38个配对，差值为0.013（[0.000, 0.039]），其中1题G3较高、37题相同、没有观察到G3较低。区间以可用配对为条件，不校正多重比较，也不是预注册检验。由于G2/G3缺失较多，不能据可用配对区间推断全60题都存在正向收益；表10的全题缺失界均跨零。")
    replace_paragraph(doc, "总分变化主要来自哪些维度", "在可配对的供应商输出中，G2−G1的公司宏维度差异为：公司事实−0.009、法规依据+0.237、引用支持0.000、不确定性0.000，总分差+0.228。因此观察到的总分差主要集中在法规评分维度，不能改写为各方面质量普遍提高，也不能把法规评分变化解释为时点判定正确率。")
    replace_paragraph(doc, "386条回答共", "过滤后的355条供应商回答共1420对逐维度评分，精确一致率为86.20%（公司事实）、57.75%（法规依据）、97.46%（引用支持）和100.00%（不确定性）。法规项二次加权κ约0.215；人员2的法规分相对人员1平均低0.417，显示评分尺度或判断存在差异。均分能够对称保留两人的信息，但不能将这种差异自动变成共识。")
    replace_paragraph(doc, "引用与不确定性精确一致率分别为", "引用支持的精确一致率为97.46%，κ为0；不确定性评分在355条回答上两位人员均给出2分，故κ因恒定边际不可估计。高一致率伴随近满分分布，不能据此宣称逐论断测量可靠。本文不选择“方法完全相同”或“量表缺乏分辨率”中的任一种作为已证实解释。")
    replace_paragraph(doc, "新增锐界保留了缺失配对中的已知一侧", "排除会话补充回答后，G2−G1只有41/60个可配对问题。利用每题可观察的一侧并对缺失评分按0至8取界，其全题均差范围为[-1.508, 1.292]；G3−G2为[-1.567, 1.767]。两者均跨零，因此当前可用样本的正向公司聚类区间不能外推为全计划问题上的确定优势。此类范围是基于量表端点的代数缺失界，不是置信区间。逐公司删一范围与缺失界见表10和图7。")

    replace_paragraph(doc, "保利发展分红案例显示了角色提示下", "作为一项事后选取的供应商生成案例，立讯精密审计委员会问题在G1、G2、G3下的双人均分分别为7.0、8.0、8.0。G3输出依据年报报告成员、会议次数和履职情况，同时指出现有年报未逐一说明成员董事类别及召集人身份，且法规文本时点未核验，故不能仅据该回答认定符合规范。该单例用于展示输出与评分边界，不是主效应的独立证据。")
    replace_paragraph(doc, "美的审计委员会题三个条件", "在该案例中，G2与G3使用逐字相同的检索上下文且评分相同；它说明角色提示没有在此题增加评分，不能据一个案例推断提示普遍无效。案例和评分均来自主协议下的供应商输出，已排除会话补充答案。")

    replace_paragraph(doc, "图4 原主结果与人员敏感性", "图4 三个主条件的供应商输出评分与评分者差异。均值及点对点95%区间仅基于各条件可用输出，区间按公司聚类。")
    replace_paragraph(doc, "图5 维度配对效应。10000次公司聚类bootstrap", "图5 主比较的公司宏维度配对差。区间为20,000次公司聚类bootstrap点对点95%区间；仅包括已观察配对，不作多重比较校正。")
    replace_paragraph(doc, "图6 评分天花板与人员满分率", "图6 355条供应商输出上的评分天花板与评分者满分率。高一致率不代表逐论断语义支持已核验。")
    replace_paragraph(doc, "图7 缺失与逐公司删一敏感性", "图7 排除会话补充回答后的缺失与逐公司删一敏感性。观测效应区间为20,000次公司聚类bootstrap点对点95%区间；全题代数缺失界与置信区间是不同估计对象。")
    replace_paragraph(doc, "表12 事后真实案例", "表12 事后选取的供应商输出案例")

    replace_paragraph(doc, "执行核验包括：400个计划源文件身份", "执行核验包括400个计划单元、355条纳入分析的供应商输出、386条历史评分记录及其3,088个原始评分字段；其中31条会话补充回答已从重分析中排除。原配置120个上下文逐字一致，参考证据123条参考来源位置均可在索引中定位。评分者在正文仅以人员1、人员2表示；历史评分和失败记录保留，新增分析不回溯称为预注册。")

    update_tables(doc)
    doc.save(OUT)
    replace_embedded_media(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
