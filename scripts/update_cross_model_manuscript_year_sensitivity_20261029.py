"""Create a new paper candidate that reports the saved cross-model results by year."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_统一主文_跨模型与Agent对照候选稿_待终审.docx"
ANALYSIS = ROOT / "results/runs/qwen_cross_model_year_sensitivity_20261029_01/year_stratified_summary.json"
OUTPUT = ROOT / "results/paper/condensed_candidate_20261028_01/年报与法规双源RAG_统一主文_跨模型与Agent年度分层复算候选稿_待终审.docx"
RECEIPT = ROOT / "results/runs/qwen_cross_model_year_sensitivity_20261029_01/manuscript_update_receipt.json"
ANCHOR = "补充开展跨模型族、跨年度的小样本探索："


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main() -> None:
    for path in (SOURCE, ANALYSIS):
        if not path.is_file():
            raise FileNotFoundError(path)
    if OUTPUT.exists() or RECEIPT.exists():
        raise SystemExit("Refusing to overwrite an existing manuscript candidate or receipt.")

    report = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    strata = {
        (item["reference"], item["year"], item["model"]): item
        for item in report["strata"]
    }
    corrected = "post_output_erratum"
    expected_keys = [(corrected, year, model) for year in (2023, 2024)
                     for model in ("qwen06", "k3_single")]
    if any(key not in strata for key in expected_keys):
        raise ValueError("Year-stratified data are incomplete")

    q23, k23 = (strata[(corrected, 2023, model)] for model in ("qwen06", "k3_single"))
    q24, k24 = (strata[(corrected, 2024, model)] for model in ("qwen06", "k3_single"))
    def fmt(row, metric):
        return f'{row[f"{metric}_macro_nonempty"]:.3f}'

    sentence = (
        " 按报告年度分层（每年4题，均来自相同4家公司），勘误后参考下2023年Qwen/K3的"
        f"成员宏F1为{fmt(q23, 'members_f1')}/{fmt(k23, 'members_f1')}、"
        f"日期宏F1为{fmt(q23, 'dates_f1')}/{fmt(k23, 'dates_f1')}、"
        f"会议次数精确匹配{int(q23['meeting_count_exact_macro_nonempty'] * 4)}/4与"
        f"{int(k23['meeting_count_exact_macro_nonempty'] * 4)}/4、三字段全匹配"
        f"{int(q23['all_three_exact_macro_nonempty'] * 4)}/4与"
        f"{int(k23['all_three_exact_macro_nonempty'] * 4)}/4；2024年相应为"
        f"{fmt(q24, 'members_f1')}/{fmt(k24, 'members_f1')}、"
        f"{fmt(q24, 'dates_f1')}/{fmt(k24, 'dates_f1')}、"
        f"{int(q24['meeting_count_exact_macro_nonempty'] * 4)}/4与"
        f"{int(k24['meeting_count_exact_macro_nonempty'] * 4)}/4、"
        f"{int(q24['all_three_exact_macro_nonempty'] * 4)}/4与"
        f"{int(k24['all_three_exact_macro_nonempty'] * 4)}/4。"
        f"两年度Qwen/K3的严格结构合规和引用列表合规均为0/4与4/4；"
        f"Qwen引用列表中机械解析到的格式异常字符串项分别为"
        f"{q23['malformed_citation_list_entry_count']}和"
        f"{q24['malformed_citation_list_entry_count']}（不含标量引用值）。"
        "每年仅4题且为相同公司，故此为描述性分层，不是独立时间迁移验证，也不作显著性推断。"
    )

    document = Document(SOURCE)
    matches = [p for p in document.paragraphs if p.text.startswith(ANCHOR)]
    if len(matches) != 1:
        raise ValueError(f"Expected one cross-model paragraph, found {len(matches)}")
    matches[0].add_run(sentence)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    receipt = {
        "source_docx": str(SOURCE.relative_to(ROOT)),
        "source_docx_sha256": sha256(SOURCE),
        "analysis_json": str(ANALYSIS.relative_to(ROOT)),
        "analysis_json_sha256": sha256(ANALYSIS),
        "output_docx": str(OUTPUT.relative_to(ROOT)),
        "output_docx_sha256": sha256(OUTPUT),
        "anchor": ANCHOR,
        "inserted_text": sentence,
        "source_paragraphs": len(Document(SOURCE).paragraphs),
        "output_paragraphs": len(Document(OUTPUT).paragraphs),
        "no_new_model_calls": True,
        "no_new_human_annotations": True,
    }
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == "__main__":
    main()
