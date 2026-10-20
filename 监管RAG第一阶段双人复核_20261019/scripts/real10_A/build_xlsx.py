# -*- coding: utf-8 -*-
"""复核人A(A-7402) real10 复核表构建：三页填写 + 回答级判断"""
import json, os, shutil, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

SRC='review_work/R1/real10/复核表.xlsx'
DST='submissions/reviewer_A-7402/real10/复核表.xlsx'
os.makedirs(os.path.dirname(DST), exist_ok=True)

RID='A-7402'; DATE='2026-10-19'
BLIND='独立填写；提交前未查看另一复核人表格或判断；未使用包外材料'

# 回答级判断: (可回答性, 判断理由, 不确定原因)
LEVEL={
'D01':('可回答',
 '28条论断核对完毕：26条获保存证据联合支持；2条（独董个人出席次数/方式、独董专门会议机制）因保存证据未含且属对年报全文的否定性陈述判无法判断，回答以“无法核验/无法确认”收场，与证据状态一致；引文与A1/A3/A4/A5及R1-R5核对一致（异议栏目勾选原文为“□适用√不适用”，回答作“：不适用”文本压缩）；A2（封面块）在正文无引用去向，已按规范登记未关联、支持填否；法规均标注时点未核验，回答未作强制适用。',
 '法规时点适用性未经原件核验；2条否定性陈述涉及年报全文范围，待第二阶段核原件。'),
'D02':('部分可回答',
 '14条论断均获保存证据联合支持；回答将“无关联交易/无违规担保/无资金占用”限定为公司单方披露证据，未作独立核实；担保全貌（发生额、余额、对象）与决策程序披露在证据中缺失，回答如实说明无法评价；章程指引2019/2023/2025各版本关联交易与独董过半数同意要求引用正确；法规时点未核验已声明，未把未披露推断为不存在。',
 '担保全貌与决策程序是否在年报其他位置披露，待第二阶段核原件；法规版本适用性未确认。'),
'D03':('部分可回答',
 '32条论断中30条获支持、2条判无法判断（“本案未发行优先股”前提无保存证据；内部审计机构设置属年报范围否定性陈述）；发现引文瑕疵：C21引号内漏“集团”二字（A3原文为“了解集团与确定合同中履约义务完成时间相关…”）、C27与C29引号内为压缩概括（实质一致）、C24“标准无保留意见”较A1原文“标准的无保留意见”少“的”字；C26正文写“A3显示”而句末标注A2（该句实在A2物理页59，标注A2正确，文字指向有误）；R4为2025年文本晚于报告期，未强制适用。',
 '2条无法判断项待第二阶段核原件；引文压缩瑕疵不影响实质支持，但影响逐字引用准确性评价。'),
'D04':('部分可回答',
 '26条论断均获支持；回答如实指出股东大会表在第四次决议内容处截断、缺全年会议完整清单/出席比例/表决明细；正确排除R2北交所发行注册办法的主题误用；如实揭示1个非财务报告内控重大缺陷（印章管理）及其整改，但不据此认定会议程序缺陷，分寸恰当；R5为2025年文本晚于报告期已声明。',
 '会议完整清单与述职履行情况待第二阶段核原件；法规时点未确认。'),
'D05':('可回答',
 '24条论断均获支持；财务报表审计与内控审计均为标准无保留意见、评价范围覆盖合并口径100%、自评与鉴证一致、2025-04-29巨潮披露日期等逐项核对一致；C07句中内控审计报告披露日期实际载于A2（物理页55）而该句仅标注A3，标注不完整已按实际去向登记；“合理保证非绝对保证”等边界表述准确；回答限定在年报披露口径内作结论，未越界。',
 'A2表头“审议意见段”疑为“审计意见段”文本识别误差，待第二阶段核原件；法规时点未确认。'),
'D06':('部分可回答',
 '30条论断均获支持；回答核实并如实指出3月20日议案名称《关于提请召开2024年年度股东大会的议案》与4月11日实际召开“2023年年度股东大会”的称谓不一致，作开放定性、未直接认定违法，处理符合规范；邵振安股东大会出席0次因系报告期补选、任职时点不明，未作违规评价；信披办法2025-02-28版本对2024年度的适用性未确认已声明；会议通知日期、表决票型缺失如实说明。',
 '议案名称不一致的性质（文本瑕疵或程序问题）待第二阶段核原件；任职时点与通知日期不明。'),
'D07':('部分可回答',
 '17条论断均获支持；回答将“无处罚/无事故”限定为公司陈述未扩展为外部事实；指出A2仅为页眉、A4仅为封面，风险因素章节正文不在证据内，无法评价重大风险披露充分性；正确识别R1（期货衍生品）、R2（融资担保）、R5（股权激励）与主题不匹配，R3仅为合规手册索引而非独立强制规范；内控审计意见未被外延解释为环境安全事项，边界清楚。',
 '风险因素章节内容、安全节截断处后续内容待第二阶段核原件；法规时点未确认。'),
'D08':('可回答',
 '14条论断均获支持；分红金额经算术复核一致（782,704,094÷10×0.8=62,616,327.52元）；预案核心要素在A1/A2/A3三处口径一致；“尚需经2024年度股东大会审议”的程序状态明确，未断言已批准或已实施；章程指引2019/2023文本条款引用一致，版本与生效适用性未确认已声明。',
 '章程指引版本适用性与股东大会最终表决结果待第二阶段核原件。'),
'D09':('部分可回答',
 '31条论断均获支持；回答自评“部分支持”与证据状态一致；准确保留A5“财务报告重大缺陷数量”行页面截断问题（经核片段确以“非财务报告重大缺陷数量”起始）；正确识别R5为科创板规则且标注2026-08-01，与深市主板不匹配，点名排除；R1治理准则标注2026-01-01晚于报告期不作确定准据；发现引文压缩瑕疵3处（C08、C23、C27，实质一致）及“董监高”概括略宽于R4原文“董事、高级管理人员”；A3（封面）在正文无方括号引用去向（“年报披露于2025-03”一句未标注来源），已登记未关联、支持填否。',
 '法规时点与适用版本未确认；引文压缩不影响实质支持但影响逐字引用准确性评价。'),
'D10':('部分可回答',
 '27条论断均获支持；审计委员会成员、5个会议日期、议题清单、审计意见类型/签署日期/文号/签字会计师及签字盖章页页码逐项核对一致；7次会议仅5个日期、无法核验季度分布如实说明；R1章程指引标注2025-03-28晚于报告期，不作追溯性合规判断；C07/C08句末合并标注[R1][R3]，已按论断-来源实际对应关系分别登记（R3支撑第二十六条句、R1支撑第一百三十五条句，交叉角色填否）；正确区分2019/2023/2025三版章程指引非标意见说明条款且本案为标准无保留意见未触发。',
 '会议出席人数、表决结果、季度分布及期后审议情况待第二阶段核原件；法规时点未确认。'),
}

ORDER=['D01','D02','D03','D04','D05','D06','D07','D08','D09','D10']

shutil.copy(SRC, DST)
wb=openpyxl.load_workbook(DST)

# --- 回答级判断 ---
ws=wb['回答级判断']
for r in range(2,12):
    q=ws.cell(r,1).value
    ans_level=LEVEL[q]
    ws.cell(r,3).value=RID
    ws.cell(r,4).value=ans_level[0]
    ws.cell(r,5).value='带保留回答'
    ws.cell(r,6).value='是'
    ws.cell(r,7).value='是'
    ws.cell(r,8).value=None
    ws.cell(r,9).value='保留'
    ws.cell(r,10).value=ans_level[1]
    ws.cell(r,11).value=ans_level[2]
    ws.cell(r,12).value=BLIND
    ws.cell(r,13).value=DATE

# --- 论断划分与支持 ---
ws=wb['论断划分与支持']
def fill_block(ws, q, rows_data, ncols, writer):
    # 找到该题块行范围
    rows=[r for r in range(2,ws.max_row+1) if ws.cell(r,1).value==q]
    need=len(rows_data); have=len(rows)
    if need>have:
        ws.insert_rows(rows[-1]+1, need-have)
        for i in range(need-have):
            ws.cell(rows[-1]+1+i,1).value=q
        rows=[r for r in range(2,ws.max_row+1) if ws.cell(r,1).value==q]
    for r,data in zip(rows,rows_data):
        writer(r,data)
    return len(rows)

answers={q:open(f'review_work/R1/real10/阅读材料/{q}_回答原文.txt',encoding='utf-8').read() for q in ORDER}
# 从后往前填，避免 insert_rows 打乱后续块定位
for q in reversed(ORDER):
    d=json.load(open(f'analysis_A/{q}.json',encoding='utf-8'))
    for c in d['claims']:
        cnt=answers[q].count(c[1])
        assert cnt==1,(q,c[0],cnt)
    def writer(r,c,ws=ws):
        ws.cell(r,2).value=c[0]      # 论断编号
        ws.cell(r,3).value=c[1]      # 回答原文引句
        ws.cell(r,4).value=None      # 同句出现序号（仅出现一次，留空）
        ws.cell(r,5).value=c[2]      # 论断类型
        ws.cell(r,6).value=c[3]      # 需要年报证据
        ws.cell(r,7).value=c[4]      # 需要法规证据
        ws.cell(r,8).value=c[5]      # 保存证据联合支持
        ws.cell(r,9).value=c[6]      # 支持证据ID
        ws.cell(r,10).value=c[7]     # 需要时点判断
        ws.cell(r,11).value=c[8]     # 时点适用性
        ws.cell(r,12).value=c[9]     # 判断理由
        ws.cell(r,13).value=c[10]    # 不确定原因
    fill_block(ws,q,d['claims'],13,writer)

# --- 引用关联 ---
ws=wb['引用关联']
for q in reversed(ORDER):
    d=json.load(open(f'analysis_A/{q}.json',encoding='utf-8'))
    def writer(r,l,ws=ws):
        ws.cell(r,2).value=l[0]      # 论断编号或未关联
        ws.cell(r,3).value=l[1]      # 引用ID
        ws.cell(r,4).value=l[2]      # 支持其来源角色
        ws.cell(r,5).value=l[3]      # 时点适用性
        ws.cell(r,6).value=l[4]      # 判断理由
        ws.cell(r,7).value=l[5] or None  # 不确定原因
    fill_block(ws,q,d['links'],7,writer)

wb.save(DST)
print('saved:',DST)

# 校验输出
wb2=openpyxl.load_workbook(DST)
ws=wb2['论断划分与支持']
from collections import Counter
cnt=Counter()
for r in range(2,ws.max_row+1):
    if ws.cell(r,2).value: cnt[ws.cell(r,1).value]+=1
print('claims filled:',dict(cnt))
ws=wb2['引用关联']
cnt=Counter()
for r in range(2,ws.max_row+1):
    if ws.cell(r,3).value: cnt[ws.cell(r,1).value]+=1
print('links filled:',dict(cnt))
ws=wb2['回答级判断']
print('level rows:',[(ws.cell(r,1).value,ws.cell(r,4).value,ws.cell(r,9).value) for r in range(2,12)])
