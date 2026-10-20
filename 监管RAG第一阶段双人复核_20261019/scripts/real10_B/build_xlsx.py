# -*- coding: utf-8 -*-
"""复核人B(B-9136) real10 复核表构建：三页填写 + 回答级判断"""
import json, os, shutil, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

SRC='review_work/R2/real10/复核表.xlsx'
DST='submissions/reviewer_B-9136/real10/复核表.xlsx'
os.makedirs(os.path.dirname(DST), exist_ok=True)

RID='B-9136'; DATE='2026-10-19'
BLIND='由同一AI顺序扮演两名复核人：A-7402的提交表先于本表完成，填写本表时能见到A表内容，无法声明完全盲态；本表判断系从R2包证据独立重新推导，未使用包外材料；与A判断分歧处如实保留，待协调者裁定'

LEVEL={
'D01':('可回答',
 '25条论断核对：23条获保存证据联合支持；K21（独董个人出席次数/方式未披露）、K22（独董专门会议机制）属年报范围否定性陈述判无法判断，回答以“无法核验/无法确认”收场，与证据状态一致；委员会构成、会议次数（2/5/1/6）、议题清单与A1/A3/A4/A5逐项一致；A2（封面块）无正文引用去向，已登记未关联、支持填否；R4治理准则标注2026-01-01晚于报告期未作强制基准；R1-R5时点未核验已声明。',
 '2条否定性陈述及法规时点适用性待第二阶段核原件。'),
'D02':('部分可回答',
 '11条论断均获支持；关联交易六项否定、无非经营性占用、无违规担保与A1/A2一致；回答正确限定“违规否定≠无任何担保”、决策程序披露缺失无法评价、法规仅作框架参考；章程指引三版本条款对应（R5/R3为第一百一十条、R4为第一百一十三条、R2第一百三十一条）核对一致；四点评价限制与证据状态相符。',
 '担保全貌与决策程序是否另载于年报其他位置、法规版本适用性，待第二阶段核原件。'),
'D03':('部分可回答',
 '27条论断中25条获支持、2条判无法判断（K08“本案未发行优先股”前提保存证据未载；K26内部审计机构设置属年报范围否定性陈述）；引文瑕疵三处：K16引文漏“集团”二字（A3原文“了解集团与确定合同中…”）、K18/K22为压缩引文（实质一致）、K21句内称“A3显示”而内容实在A2（句末标注A2正确，行文指向有误）；未提供内控审计报告正文、无法独立核验无保留结论边界的说明正确。',
 '2条无法判断项与引文准确性问题待第二阶段核原件。'),
'D04':('部分可回答',
 '21条论断均获支持；四次股东大会日期、次日刊登、全部通过与A2一致，第四次决议栏截断属实；章程修订/独董制度/关联交易三例证的程序衔接与A3一致；R2北交所发行注册条款排除正确；印章管理重大缺陷及整改与A5一致，未扩大为会议程序缺陷；述职履行保留、披露时限不作合规认定，分寸恰当。',
 '全年会议完整清单、述职履行情况与届时披露时限规则待第二阶段核原件。'),
'D05':('可回答',
 '16条论断均获支持；审计意见类型/机构/文号/签署日/签字会计师与A4一致；内控审计意见段与表格四项载明与A2一致；缺陷四项0与100%比例与A3一致；K05句中“内控审计报告2025-04-29巨潮披露”要素实在A2而句仅标A3，标注不完整已记录（事实本身两文件日期相同，结论不受影响）；R1证券公司条例、R5优先股条款排除适用正确；合理保证非绝对的边界表述准确。',
 '法规时点未确认；K05标注不完整不影响事实结论。'),
'D06':('部分可回答',
 '21条论断均获支持；5次董事会日期、出席表、议案清单与A1一致；K18会议称谓不一致属实（A1议案名《关于提请召开2024年年度股东大会的议案》与A4实际“2023年年度股东大会”），回答如实发现且未直接认定程序违法，处理正确；邵振安出席0次因报告期补选未作违规评价正确；专委会前置审议日期对照（3-13/4-7/8-9/10-12均早于董事会）核实成立；R5信披办法2025-02-28适用性未确认已声明。',
 '称谓不一致的性质、任职生效时点、通知日期与表决票型待第二阶段核原件。'),
'D07':('部分可回答',
 '13条论断均获支持；三基地备案时点、投入税额、处罚/事故“无”栏与A1一致；安全生产节在“定期对公司内重点区域进行认”处截断属实；A2仅页眉（26字）、A4仅封面（63字），风险章节缺失判断正确；R1期货衍生品、R2融资担保、R5股权激励（2026-08-01）主题不匹配排除正确；R3为合规手册索引而非强制原文的定性正确；内控审计意见未外延解释为环境安全事项，边界正确。',
 '风险因素章节内容与安全生产节截断后续内容待第二阶段核原件。'),
'D08':('可回答',
 '8条论断均获支持；分配预案要素在A1/A2/A3三处一致，算术复核782,704,094÷10×0.8=62,616,327.52元成立；“尚需经2024年度股东大会审议”状态明确，未断言已批准或实施；K02合并标注中A3不含金额62,616,327.52元要素已记录；R2/R3第一百零七条（含超授权提交股东大会）、R1第一百五十六条引用一致；回报规划与章程一致性因全文未提供不作核验。',
 '章程指引版本适用性与股东大会最终表决结果待第二阶段核原件。'),
'D09':('部分可回答',
 '29条论断均获支持；会议次数5/10/9在A1/A2双重印证；A3（封面块72字）无正文引用去向已登记未关联、支持填否；K18“财务报告重大缺陷数量”行截断属实（A5块自“非财务报告重大缺陷数量”行起始）；R5科创板发行注册规则（2026-08-01）板块不匹配排除正确；R1治理准则2026-01-01晚于报告期不作确定准据；R3/R4信披办法（2025-02-28）跨年报披露时点（2025-03-29）的版本未确定分析成立。',
 '法规适用版本、独董比例与审计委员会构成等条款缺项待第二阶段核原件。'),
'D10':('部分可回答',
 '24条论断均获支持；审委会成员（潘同文/张永龙/王天广）与身份（A1/A3）、7次会议仅5日期、议题清单、差错更正议案、审计意见要素（A4物理69页）与签字盖章页（A5物理72页）逐项一致；R3第二十六条（含每季度至少一次、三分之二出席）与R1第一百三十五条的归属准确；三版章程指引非标说明条款（R5/R4第一百零八条、R2第一百一十一条）引用正确，本案标准无保留未触发；K14期后审议不作“未审议”认定，符合规范。',
 '出席人数、表决结果、季度分布、期后审议与差错更正性质待第二阶段核原件。'),
}

ORDER=['D01','D02','D03','D04','D05','D06','D07','D08','D09','D10']

shutil.copy(SRC, DST)
wb=openpyxl.load_workbook(DST)

ws=wb['回答级判断']
for r in range(2,12):
    q=ws.cell(r,1).value
    lv=LEVEL[q]
    ws.cell(r,3).value=RID
    ws.cell(r,4).value=lv[0]
    ws.cell(r,5).value='带保留回答'
    ws.cell(r,6).value='是'
    ws.cell(r,7).value='是'
    ws.cell(r,8).value=None
    ws.cell(r,9).value='保留'
    ws.cell(r,10).value=lv[1]
    ws.cell(r,11).value=lv[2]
    ws.cell(r,12).value=BLIND
    ws.cell(r,13).value=DATE

def fill_block(ws, q, rows_data, writer):
    rows=[r for r in range(2,ws.max_row+1) if ws.cell(r,1).value==q]
    need=len(rows_data); have=len(rows)
    if need>have:
        ws.insert_rows(rows[-1]+1, need-have)
        for i in range(need-have):
            ws.cell(rows[-1]+1+i,1).value=q
        rows=[r for r in range(2,ws.max_row+1) if ws.cell(r,1).value==q]
    for r,data in zip(rows,rows_data):
        writer(r,data)

answers={q:open(f'review_work/R2/real10/阅读材料/{q}_回答原文.txt',encoding='utf-8').read() for q in ORDER}

ws=wb['论断划分与支持']
for q in reversed(ORDER):
    d=json.load(open(f'analysis_B/{q}.json',encoding='utf-8'))
    for c in d['claims']:
        cnt=answers[q].count(c[1])
        assert cnt==1,(q,c[0],cnt)
    def writer(r,c,ws=ws):
        ws.cell(r,2).value=c[0]; ws.cell(r,3).value=c[1]; ws.cell(r,4).value=None
        ws.cell(r,5).value=c[2]; ws.cell(r,6).value=c[3]; ws.cell(r,7).value=c[4]
        ws.cell(r,8).value=c[5]; ws.cell(r,9).value=c[6]; ws.cell(r,10).value=c[7]
        ws.cell(r,11).value=c[8]; ws.cell(r,12).value=c[9]; ws.cell(r,13).value=c[10] or None
    fill_block(ws,q,d['claims'],writer)

ws=wb['引用关联']
for q in reversed(ORDER):
    d=json.load(open(f'analysis_B/{q}.json',encoding='utf-8'))
    def writer(r,l,ws=ws):
        ws.cell(r,2).value=l[0]; ws.cell(r,3).value=l[1]; ws.cell(r,4).value=l[2]
        ws.cell(r,5).value=l[3]; ws.cell(r,6).value=l[4]; ws.cell(r,7).value=l[5] or None
    fill_block(ws,q,d['links'],writer)

wb.save(DST)
print('saved:',DST)

wb2=openpyxl.load_workbook(DST)
from collections import Counter
ws=wb2['论断划分与支持']
cnt=Counter()
for r in range(2,ws.max_row+1):
    if ws.cell(r,2).value: cnt[ws.cell(r,1).value]+=1
print('claims filled:',dict(cnt))
ws=wb2['引用关联']
cnt=Counter()
for r in range(2,ws.max_row+1):
    if ws.cell(r,3).value: cnt[ws.cell(r,1).value]+=1
print('links filled:',dict(cnt))
ws=wb2['回答级判断']
print('level rows:',[(ws.cell(r,1).value,ws.cell(r,4).value,ws.cell(r,9).value) for r in range(2,12)])
